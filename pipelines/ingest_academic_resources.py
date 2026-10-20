import sys
from pathlib import Path
_ROOT = Path(__file__).resolve().parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))
_CORE = _ROOT / "core"
if str(_CORE) not in sys.path:
    sys.path.insert(0, str(_CORE))

"""
法學名師與課輔講義資源整合管道 (Ingest Academic Resources)
自動讀取 G:\\我的雲端硬碟\\資源 內各法學考科講義與簡報（過濾近代物理與DLCV等非關科目），
提取名師體系、爭點肯否二說、課輔實例題與解答，蒸餾為標準法學 SFT 訓練資料集。
"""
import os
import sys
import io
import re
import json
import hashlib
from pathlib import Path
from typing import List, Dict, Any, Tuple

# Windows 終端編碼防護
if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

import pdfplumber
try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

from config import DATA_DIR

RESOURCES_DIR = Path(os.getenv("LEGAL_RESOURCES_DIR", r"G:\我的雲端硬碟\資源"))
if not RESOURCES_DIR.exists():
    for fallback_candidate in [DATA_DIR.parent / "資源", Path(r"G:\My Drive\資源"), Path(r"G:\我的雲端硬碟\資源")]:
        if fallback_candidate.exists():
            RESOURCES_DIR = fallback_candidate
            break

TRAINING_DIR = Path(DATA_DIR) / "training_data"
RAW_EXAMS_DIR = Path(DATA_DIR) / "raw_exams"
OUT_SFT_FILE = TRAINING_DIR / "law_exam_coursework_sft.jsonl"
OUT_VERIFIED_FILE = TRAINING_DIR / "law_exam_verified_sft.jsonl"
COURSEWORK_INDEX_FILE = RAW_EXAMS_DIR / "academic_coursework_index.json"

LEGAL_SUBJECT_MAP = {
    "民總": "民法總則",
    "債二": "民法債編總論",
    "債各": "民法債編各論",
    "身分法": "民法親屬繼承（身分法）",
    "刑總": "刑法總則",
    "刑分": "刑法分則",
    "憲法": "憲法與人權保障",
    "民訴": "民事訴訟法",
    "商事法": "公司法與證券交易法",
    "保險法": "商事法與保險法",
    "專利法": "智慧財產權與專利法",
    "法理學錄音": "基礎法學與法理學",
    "雜項": "國民法官法與實務專題"
}

# 嚴格排除非法律領域與系統檔案
EXCLUDE_DIRS = {"近代物理", "DLCV", "training", ".git", ".idea"}

def compute_hash(text: str) -> str:
    return hashlib.sha256(text.strip().encode("utf-8", errors="ignore")).hexdigest()

def extract_text_from_pdf(pdf_path: Path) -> str:
    """從 PDF 檔擷取完整純文字 (優先使用超高速 C++ pypdfium2，只有在 pdfium 異常時備援 pdfplumber)"""
    if pdfium:
        try:
            doc = pdfium.PdfDocument(str(pdf_path))
            try:
                text_chunks = []
                for page in doc:
                    textpage = page.get_textpage()
                    t = textpage.get_text_range()
                    if t and t.strip():
                        text_chunks.append(t.strip())
                return "\n\n".join(text_chunks)
            finally:
                doc.close()
        except Exception:
            pass

    text_chunks = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t and t.strip():
                    text_chunks.append(t.strip())
    except Exception as e:
        print(f"    [PDF讀取警示] {pdf_path.name}: {e}")
    return "\n\n".join(text_chunks)

def extract_text_from_pptx(pptx_path: Path) -> str:
    """從 PPTX 檔擷取簡報文字"""
    if not Presentation:
        return ""
    text_chunks = []
    try:
        prs = Presentation(pptx_path)
        for slide in prs.slides:
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                text_chunks.append("\n".join(slide_texts))
    except Exception as e:
        print(f"    [PPTX讀取警示] {pptx_path.name}: {e}")
    return "\n\n".join(text_chunks)

def distill_sft_samples(raw_text: str, subject: str, doc_name: str) -> List[Dict[str, Any]]:
    """
    將長篇講義文字依章節、案例與爭點，蒸餾為高品質 SFT (Instruction / Input / Output) 訓練樣本
    """
    samples = []
    if not raw_text or len(raw_text.strip()) < 100:
        return samples

    # 1. 案例與實例題萃取（課輔實例題 + 擬答）
    # 匹配特徵：[案例]、案例一、實例題演練、試問、擬答
    case_blocks = re.split(r"(?:\n|^)(?:\[案例[^\]]*\]|案例[一二三四五六七八九十\d]+|【實例題】|【案例演練】|實例題[一二三四五\d]*)(?:：|:|\s*)", raw_text)
    if len(case_blocks) > 1:
        for idx, block in enumerate(case_blocks[1:], 1):
            block_clean = block.strip()
            if len(block_clean) > 150:
                # 嘗試分割出題幹與解答
                parts = re.split(r"(?:擬答|分析|解析|解說|解答|參考擬答|檢討)(?:：|:|\n)", block_clean, maxsplit=1)
                if len(parts) == 2:
                    q_fact = parts[0].strip()
                    ans_text = parts[1].strip()
                    if len(q_fact) > 50 and len(ans_text) > 80:
                        samples.append({
                            "id": f"cw_{subject}_{compute_hash(q_fact)[:12]}",
                            "source_type": "verified_academic_coursework",
                            "instruction": f"你是一位具備多年命題閱卷經驗之台灣司法官律師考試閱卷委員。請針對以下【{subject}】名師課輔實例題，依據嚴謹三段論法（開標、大前提、實務學說見解對比、小前提涵攝）進行深度評析與提出高分示範擬答：",
                            "input": f"【案例事實與問題】：\n{q_fact[:3000]}",
                            "output": ans_text[:4000]
                        })

    # 2. 爭點與肯否兩說對比萃取
    # 匹配特徵：學說、實務、甲說、乙說、肯定說、否定說、折衷說、通說
    paragraphs = [p.strip() for p in raw_text.split("\n\n") if len(p.strip()) > 100]
    for p in paragraphs:
        has_debate = any(k in p for k in ["甲說", "乙說", "肯定說", "否定說", "通說", "有力說", "學說與實務", "兩說"])
        has_citations = any(k in p for k in ["條", "字第", "釋字", "最高法院", "判決", "裁定", "大法庭", "憲判"])
        if has_debate and has_citations and len(p) >= 120:
            lines = p.split("\n")
            title = lines[0].strip()[:60] if lines else "爭議核心"
            samples.append({
                "id": f"cw_debate_{subject}_{compute_hash(p)[:12]}",
                "source_type": "verified_academic_coursework",
                "instruction": f"請針對台灣法學界與實務就【{subject}】之重要法律爭點，深入對比「學說通說/有力說」與「實務判例/裁定判旨」之立論差異，並依據法益衡量提出嚴謹的評析結論：",
                "input": f"【探討主題】：{title}\n【講義論述焦點】：\n{p[:2000]}",
                "output": f"【爭點定位與論述架構】：\n本件涉及{subject}核心規範意旨。\n\n【論理深入剖析】：\n{p[:3500]}"
            })

    # 3. 重要法條構成要件與名師專題論述
    # 遍歷各大段落，擷取具有體系架構之論述
    for idx in range(0, len(paragraphs), 3):
        chunk = "\n\n".join(paragraphs[idx:idx + 3])
        if len(chunk) > 300 and any(k in chunk for k in ["要件", "效力", "請求權", "責任", "原則", "構成要件"]):
            first_line = chunk.split("\n")[0].strip()[:50]
            samples.append({
                "id": f"cw_doctrine_{subject}_{compute_hash(chunk)[:12]}",
                "source_type": "verified_academic_coursework",
                "instruction": f"請闡述台灣【{subject}】中關於「{first_line}」之法律規範基礎、核心構成要件、權利義務效果與相關實務判解發展：",
                "input": f"【主題研討】：{first_line}",
                "output": f"【完整體系論述】：\n{chunk[:3500]}"
            })

    return samples

def ingest_all_resources():
    print("=" * 75)
    print("  開始執行：法學名師與課輔講義資源整合與 SFT 資料蒸餾管道")
    print(f"  來源目錄：{RESOURCES_DIR}")
    print("=" * 75)

    if not RESOURCES_DIR.exists():
        print(f"[錯誤] 找不到資源目錄：{RESOURCES_DIR}")
        return

    all_samples = []
    coursework_catalog = []
    seen_hashes = set()

    # 遍歷 G:\我的雲端硬碟\資源
    for sub_dir_name in os.listdir(RESOURCES_DIR):
        sub_path = RESOURCES_DIR / sub_dir_name
        if not sub_path.is_dir() or sub_dir_name in EXCLUDE_DIRS:
            continue

        subject = LEGAL_SUBJECT_MAP.get(sub_dir_name, sub_dir_name)
        print(f"\n📂 正在處理考科：【{sub_dir_name}】 -> 標準學門：{subject}")

        # 遍歷該考科內所有子檔案（含遞迴子資料夾）
        file_count = 0
        subject_samples = 0
        for root, _, files in os.walk(sub_path):
            for file_name in files:
                ext = Path(file_name).suffix.lower()
                if ext not in [".pdf", ".pptx"]:
                    continue

                full_file_path = Path(root) / file_name
                file_count += 1
                
                # 提取文字
                if ext == ".pdf":
                    text = extract_text_from_pdf(full_file_path)
                elif ext == ".pptx":
                    text = extract_text_from_pptx(full_file_path)
                else:
                    text = ""

                if not text or len(text.strip()) < 50:
                    continue

                # 建立在線檢索目錄
                coursework_catalog.append({
                    "subject": subject,
                    "folder": sub_dir_name,
                    "filename": file_name,
                    "relative_path": str(full_file_path.relative_to(RESOURCES_DIR)),
                    "char_count": len(text),
                    "summary_preview": text[:300].replace("\n", " ")
                })

                # 蒸餾成 SFT 訓練樣本
                samples = distill_sft_samples(text, subject, file_name)
                for s in samples:
                    h = compute_hash(s["instruction"] + s["input"] + s["output"][:100])
                    if h not in seen_hashes:
                        seen_hashes.add(h)
                        all_samples.append(s)
                        subject_samples += 1

                print(f"    [{file_count:02d}] {file_name[:35]:<35} | {len(text):>6,} 字 | +{len(samples)} 樣本", flush=True)

        print(f"  └─ 共解析 {file_count} 份講義檔案，成功蒸餾出 {subject_samples} 組高品質 SFT 訓練樣本！")

    # 寫入專屬的課輔講義 SFT 檔案
    TRAINING_DIR.mkdir(parents=True, exist_ok=True)
    with open(OUT_SFT_FILE, "w", encoding="utf-8") as f:
        for item in all_samples:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")

    print("\n" + "=" * 75)
    print(f"  [蒸餾完成] 講義 SFT 樣本總計：{len(all_samples):,} 組！")
    print(f"  輸出檔案：{OUT_SFT_FILE}")
    print("=" * 75)

    # 儲存在線課輔知識庫索引
    RAW_EXAMS_DIR.mkdir(parents=True, exist_ok=True)
    with open(COURSEWORK_INDEX_FILE, "w", encoding="utf-8") as f:
        json.dump(coursework_catalog, f, ensure_ascii=False, indent=2)
    print(f"  [索引完成] 講義檢索庫索引已建立：{COURSEWORK_INDEX_FILE} (共收錄 {len(coursework_catalog)} 份講義)")

    # 合併入 verified SFT 庫（嚴格去重）
    merge_into_verified_dataset(all_samples)

def merge_into_verified_dataset(new_samples: List[Dict[str, Any]]):
    print("\n" + "=" * 75)
    print("  正在將名師講義高品質語料安全合併入主微調庫 (law_exam_verified_sft.jsonl)...")
    print("=" * 75)

    existing_hashes = set()
    total_existing = 0

    if OUT_VERIFIED_FILE.exists():
        with open(OUT_VERIFIED_FILE, "r", encoding="utf-8") as f:
            for line in f:
                line_str = line.strip()
                if line_str:
                    try:
                        item = json.loads(line_str)
                        h = compute_hash(item.get("instruction", "") + item.get("input", ""))
                        existing_hashes.add(h)
                        total_existing += 1
                    except Exception:
                        pass

    added_count = 0
    with open(OUT_VERIFIED_FILE, "a", encoding="utf-8") as f:
        for s in new_samples:
            h = compute_hash(s.get("instruction", "") + s.get("input", ""))
            if h not in existing_hashes:
                existing_hashes.add(h)
                f.write(json.dumps(s, ensure_ascii=False) + "\n")
                added_count += 1

    total_verified = total_existing + added_count
    size_mb = OUT_VERIFIED_FILE.stat().st_size / (1024 * 1024)

    print(f"  既有核實樣本數: {total_existing:,}")
    print(f"  新增講義蒸餾樣本: {added_count:,}")
    print(f"  合併後總樣本數: {total_verified:,} 組！")
    print(f"  微調庫檔案大小: {size_mb:.2f} MB")
    print("=" * 75)

if __name__ == "__main__":
    ingest_all_resources()
